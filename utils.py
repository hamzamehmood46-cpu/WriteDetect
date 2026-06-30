import re
import nltk
import re
import os
import docx
import PyPDF2
import shap
import pandas as pd
from fpdf import FPDF
from datetime import datetime
from nltk.corpus import stopwords, wordnet
from nltk.stem import WordNetLemmatizer
from nltk.tokenize import word_tokenize
from lime.lime_text import LimeTextExplainer


def text_preprocessing(text_input):
    stop_words = set(stopwords.words("english"))

    def clean_text(text):
        text = text.lower()
        text = re.sub(r"[^\w\s]", "", text)
        text = re.sub(r"\b[\d\w]*\d[\d\w]*\b", "", text)
        text = text.strip()
        tokens = text.split()
        tokens = [
            word
            for word in tokens
            if word not in stop_words and word.isalpha() and len(word) > 1
        ]
        return " ".join(tokens)

    if isinstance(text_input, pd.Series):
        return text_input.apply(clean_text)

    elif isinstance(text_input, str):
        return clean_text(text_input)

    elif isinstance(text_input, pd.DataFrame):
        if "essay" in text_input.columns:
            return text_input["essay"].apply(clean_text)
        else:
            raise ValueError("Expected a DataFrame with a column named 'essay'.")

    else:
        raise TypeError(
            "Input must be a string, pandas Series, or DataFrame with 'essay' column."
        )


def get_wordnet_pos(treebank_tag):
    if treebank_tag.startswith("J"):
        return wordnet.ADJ
    elif treebank_tag.startswith("V"):
        return wordnet.VERB
    elif treebank_tag.startswith("N"):
        return wordnet.NOUN
    elif treebank_tag.startswith("R"):
        return wordnet.ADV
    else:
        return wordnet.NOUN


def lemmatize_text(text, lemmatize=True):
    if not lemmatize:
        return text

    tokens = nltk.word_tokenize(text)
    pos_tags = nltk.pos_tag(tokens)
    lemmatizer = WordNetLemmatizer()
    lemmatized_tokens = [
        lemmatizer.lemmatize(word, get_wordnet_pos(tag)) for word, tag in pos_tags
    ]
    return " ".join(lemmatized_tokens)


def apply_lemmatization(series):
    return series.apply(lemmatize_text)


# Inference Explanation Using LIME
def explain_prediction_with_lime(pipeline, text):
    class_names = ["Human Written", "AI Written"]

    # Extract vectorizer and classifier
    vectorizer = pipeline.named_steps["vectorizer"]
    classifier = pipeline.named_steps["model"]

    def predict_proba(texts):
        X = vectorizer.transform(texts)
        return classifier.predict_proba(X)

    explainer = LimeTextExplainer(class_names=class_names)

    explanation = explainer.explain_instance(text, predict_proba, num_features=10)

    return explanation


# Inference Explanation Using SHAP
def explain_prediction_with_shap(pipeline, text):
    vectorizer = pipeline.named_steps["vectorizer"]
    classifier = pipeline.named_steps["model"]

    # Vectorize the input text
    X = vectorizer.transform([text])

    def predict_proba_shap(X_array):
        return classifier.predict_proba(X_array)

    background_data = shap.utils.sample(X, nsamples=1)

    # Initialize KernelExplainer
    explainer = shap.KernelExplainer(predict_proba_shap, background_data)

    shap_values = explainer.shap_values(X, nsamples=100)

    return shap_values, explainer, X


def remove_emojis(text):
    return re.sub(r"[^\x00-\x7F]+", "", text)


# Download Comparison Report
def generate_comparison_report_pdf_full(
    comparison_results,
    input_method,
    word_count=None,
    filename=None,
    original_text=None,
    lime_explanations=None,
):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Arial", "", 11)

    # Title
    pdf.set_font("Arial", "B", 15)
    pdf.cell(
        0,
        10,
        "AI vs Human Text Classification Model Comparison Report",
        ln=True,
        align="C",
    )
    pdf.set_font("Arial", "", 10)
    pdf.cell(0, 6, f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}", ln=True)
    pdf.ln(2)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(4)

    # Metadata
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 8, "Input Metadata", ln=True)
    pdf.set_font("Arial", "", 10)
    pdf.multi_cell(0, 6, f"- Input method : {remove_emojis(input_method)}")
    if filename:
        pdf.multi_cell(0, 6, f"- Filename     : {remove_emojis(filename)}")
    if word_count is not None:
        pdf.multi_cell(0, 6, f"- Word count   : {word_count}")
    pdf.ln(2)

    # Input Text
    if original_text:
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 8, "Provided Text", ln=True)
        pdf.set_font("Arial", "", 9)
        text_clean = remove_emojis(original_text.strip())
        max_chars = 3000
        if len(text_clean) > max_chars:
            text_clean = text_clean[:max_chars] + " ...[truncated]"
        pdf.multi_cell(0, 5, text_clean)
        pdf.ln(2)

    # Prediction Table
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 8, "Model Predictions", ln=True)
    pdf.set_font("Arial", "B", 9)
    col_w = [60, 30, 25, 38, 38]
    headers = ["Model", "Pred.", "Conf.", "Human %", "AI %"]
    for w, h in zip(col_w, headers):
        pdf.cell(w, 7, h, 1, 0, "C")
    pdf.ln()

    pdf.set_font("Arial", "", 9)
    chart_paths = []
    for res in comparison_results:
        row = [
            remove_emojis(res["Model"]),
            res["Prediction"],
            res["Confidence"],
            res["Human Written %"],
            res["AI Written %"],
        ]
        for w, cell in zip(col_w, row):
            pdf.cell(w, 7, cell, 1, 0, "C")
        pdf.ln()

    # Agreement Summary
    pdf.ln(2)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(2)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 8, "Agreement Summary", ln=True)
    preds = [r["Prediction"] for r in comparison_results]
    pdf.set_font("Arial", "", 10)
    if len(set(preds)) == 1:
        pdf.multi_cell(0, 6, f"All models agree: {preds[0]}")
    else:
        pdf.multi_cell(0, 6, "Disagreement detected:")
        for r in comparison_results:
            pdf.multi_cell(0, 6, f"   - {remove_emojis(r['Model'])}: {r['Prediction']}")

    # LIME Summaries
    if lime_explanations:
        pdf.ln(2)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(2)
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 8, "LIME Top-Features (per model)", ln=True)
        pdf.set_font("Arial", "", 9)
        for model, expl in lime_explanations.items():
            pdf.set_font("Arial", "B", 9)
            pdf.cell(0, 6, f"{remove_emojis(model)}", ln=True)
            pdf.set_font("Arial", "", 9)
            expl_sorted = sorted(
                [e for e in expl if abs(e[1]) > 1e-6],
                key=lambda t: abs(t[1]),
                reverse=True,
            )[:10]
            for word, weight in expl_sorted:
                pdf.cell(0, 5, f"   {word:<12} : {weight:+.3f}", ln=True)
            pdf.ln(1)

    os.makedirs("reports", exist_ok=True)
    output_path = os.path.join(
        "reports", f"ai_vs_human_comparison_report ({filename}).pdf"
    )

    pdf.output(output_path)
    return output_path


# Document Readers
def read_txt(file):
    return str(file.read(), "utf-8")


def read_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])


def read_pdf(file):
    reader = PyPDF2.PdfReader(file)
    return "".join(
        [page.extract_text() for page in reader.pages if page.extract_text()]
    )
